# -*- coding: utf-8 -*-
"""Build the Current IT Landscape facts (csf-gap-landscape-facts/1) from an Interview Assessment Excel export (csf-interview/3).
Reference implementation of Gap Analysis Stage 0 for the Landscape. The Gap Analysis HTML v1.1 applies the same rules.
Facts only: counts, lists, and recorded answers. No ratings, no gap statements.
Usage: python landscape_facts.py <interview.xlsx> [output.json]
Needs: pip install openpyxl"""
import json, sys
from collections import Counter, OrderedDict
from datetime import date
import openpyxl

BLOCKS = OrderedDict([
    ("B1", "Business and regulatory context"), ("B2", "Critical services to systems"), ("B3", "Application landscape"),
    ("B4", "Infrastructure and connectivity"), ("B5", "Identity landscape"), ("B6", "Existing security capability"),
    ("B7", "Third-party landscape"), ("B8", "Data landscape"), ("B9", "IT organization and operating model"),
    ("B10", "Control coverage snapshot"), ("B11", "Assessment basis"), ("B12", "Threat context"),
    ("B13", "Changes in flight"), ("B14", "Prior assessments and audits")])
COVN, OPN, STN = {"full": "Full", "part": "Partial", "none": "None"}, {"inh": "In-house", "mssp": "MSSP", "vend": "Vendor", "nobody": "Nobody"}, {"use": "In use", "bnd": "Bought, not deployed", "plan": "Planned"}
COVERAGE_FIELDS = ["C3", "C4", "C5", "C6", "C7", "C8", "C9"]
HOSTING = {"onprem": "on-premises data center", "priv": "private cloud or colocation", "cloud": "public cloud",
           "saas": "SaaS", "vendor": "vendor-hosted", "unk": "unknown hosting"}
ORIGIN = {"cots": "standard commercial software", "cotsc": "customized commercial software", "inh": "in-house developed",
          "saas": "SaaS", "out": "outsourced custom build"}

def rows(wb, name):
    if name not in wb.sheetnames:
        return []
    it = wb[name].iter_rows(values_only=True)
    head = [str(h) if h is not None else "" for h in next(it)]
    return [dict(zip(head, r)) for r in it if any(v not in (None, "") for v in r)]

def n(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None

def plural(k, word):
    return f"{k} {word}" + ("" if k == 1 else "s")

def join(items):
    items = [i for i in items if i]
    return items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1] if items else ""

def build(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    summary = {r.get("Item"): r.get("Value") for r in rows(wb, "Summary")}
    source = summary.get("Schema version", "")
    if source not in ("csf-interview/3", "csf-ig/4"):
        raise SystemExit(f"Interview export is {source!r}; csf-interview/3 or csf-ig/4 is required for the Landscape")
    systems = rows(wb, "Systems (Quick BIA)")
    prof = rows(wb, "System Profile")
    env = {r["Key"]: r for r in rows(wb, "Environment Profile")}
    sites, tools = rows(wb, "Sites"), rows(wb, "Security Tools")
    hist, chg, prior = rows(wb, "Incident History"), rows(wb, "Changes in Flight"), rows(wb, "Prior Assessments")
    vendors, obs = rows(wb, "Vendor Profile"), rows(wb, "Technical Observations")
    depts, dcf = rows(wb, "Departments"), rows(wb, "Department Conflicts")
    dcode = lambda d, f: str(d.get(f + " code") or "")
    pv = {}
    for r in prof:
        pv.setdefault(r["System"], {})[r["Field ID"]] = (r.get("Answer code") or "", r.get("Rating") or "", r.get("Answer") or "")
    crit = {s["System"]: (s.get("Criticality") or "") for s in systems}
    B = {k: [] for k in BLOCKS}
    add = lambda b, text, ev: B[b].append({"text": text, "evidence": ev})
    ans = lambda k: (env.get(k) or {}).get("Answer") or ""

    # B1 business and regulatory context
    if sites:
        staff = sum(n(s.get("Staff")) or 0 for s in sites)
        add("B1", f"The organization operates {plural(len(sites), 'site')}: {join([str(s.get('Site') or '') for s in sites])}" + (f", with about {staff} staff." if staff else "."), [s["ID"] for s in sites])
    if summary.get("Scope"):
        add("B1", f"Assessment scope: {summary['Scope']}", ["Summary:Scope"])

    # B2 critical services
    for s in systems:
        if crit[s["System"]] == "A":
            add("B2", f"{s['System']} supports {s.get('Business service') or 'an unnamed service'}, Criticality A, tolerable downtime {s.get('Tolerable downtime') or 'not stated'}.", [f"BIA:{s['System']}"])

    # B3 application landscape
    cc = Counter(crit[s["System"]] or "unrated" for s in systems)
    add("B3", f"{plural(len(systems), 'system')} listed: {cc.get('A', 0)} Criticality A, {cc.get('B', 0)} B, {cc.get('C', 0)} C" + (f", {cc['unrated']} unrated." if cc.get("unrated") else "."), ["POP:systems"])
    oc = Counter(pv.get(s["System"], {}).get("C1", ("",))[0] for s in systems)
    if any(oc):
        add("B3", "Origin: " + join([f"{v} {ORIGIN.get(k, k)}" for k, v in oc.most_common() if k]) + ".", ["POP:C1"])
    hc = Counter(s.get("Hosting code") or "" for s in systems)
    if any(hc):
        add("B3", "Hosting: " + join([f"{v} {HOSTING.get(k, k)}" for k, v in hc.most_common() if k]) + ".", ["POP:C10"])

    shadow = [(d["Department"], x.strip()) for d in depts for x in str(d.get("U1 Systems not on the IT list") or "").split(",") if x.strip()]
    if shadow:
        add("B3", f"Department interviews found {len(shadow)} {'system or app' if len(shadow) == 1 else 'systems or apps'} not on the IT list: " + join([f"{x} ({dn})" for dn, x in shadow]) + ".", sorted({f"U1:{dn}" for dn, _ in shadow}))

    # B4 infrastructure and connectivity
    for k, label in [("dc", "Primary data center"), ("dr", "Disaster recovery site"), ("seg", "Network segmentation"), ("rem", "Staff remote access"),
                     ("ws", "Workstations"), ("eol", "Workstations on end-of-life operating systems"), ("otn", "Networked medical devices and OT"),
                     ("otp", "Medical device and OT network placement"), ("bka", "Backup architecture")]:
        if ans(k) not in ("", None):
            add("B4", f"{label}: {ans(k)}.", [f"ENV:{k}"])
    exp = [s for s, f in pv.items() if f.get("C6", ("",))[0] == "direct"]
    if exp:
        add("B4", f"{plural(len(exp), 'system')} directly exposed to the internet: {join(exp)}.", [f"C6:{s}" for s in exp])

    # B5 identity landscape
    for fid, label in [("C3", "Authentication"), ("C4", "Multi-factor authentication (MFA)"), ("C5", "Administrator accounts")]:
        c = Counter(f.get(fid, ("", "", ""))[2] for f in pv.values())
        c.pop("", None)
        if c:
            add("B5", f"{label}: " + join([f"{k} ({v})" for k, v in c.most_common()]) + ".", [f"POP:{fid}"])
    if ans("shws"):
        add("B5", f"Shared workstations in operational areas: {ans('shws')}.", ["ENV:shws"])

    cred = [d for d in depts if dcode(d, "U2") in ("some", "shared")]
    if cred:
        add("B5", "Departments reporting shared accounts or passwords: " + join([f"{d['Department']} ({d.get('U2 Login and password sharing') or ''})" for d in cred]) + ".", [f"U2:{d['Department']}" for d in cred])
    stay = [d["Department"] for d in depts if dcode(d, "U3") == "stay"]
    if stay:
        add("B5", "Shared computers stay logged in all shift in " + join(stay) + ".", [f"U3:{n}" for n in stay])

    # B6 existing security capability
    for t in tools:
        add("B6", f"{t.get('Capability')}: {t.get('Product') or 'no product named'}, coverage {COVN.get(t.get('Coverage code'), 'not stated').lower()}, operated by {OPN.get(t.get('Operated by code'), 'not stated').lower()}, {STN.get(t.get('Status code'), 'status not stated').lower()}.", [t["ID"]])

    # B7 third-party landscape
    if vendors:
        v1 = lambda v: next((v[k] for k in v if k.startswith("V1 ") and not k.endswith(("rating", "code"))), "")
        vt = Counter(v1(v) or "" for v in vendors)
        add("B7", f"{plural(len(vendors), 'vendor')} with system or data access: " + join([f"{k} ({v})" for k, v in vt.most_common() if k]) + ".", ["POP:V1"])
        rem = [v["Vendor"] for v in vendors if str(v.get("V2 code") or "") == "rem"]
        if rem:
            add("B7", f"{plural(len(rem), 'vendor')} connect remotely: {join(rem)}.", [f"V2:{v}" for v in rem])
        ca = [v["Vendor"] for v in vendors if str(v.get("V7 code") or "") == "A"]
        if ca:
            add("B7", f"Criticality A vendors: {join(ca)}.", [f"V7:{v}" for v in ca])

    # B8 data landscape
    spd = [s for s, f in pv.items() if f.get("C11", ("",))[0] == "spd"]
    pd_ = [s for s, f in pv.items() if f.get("C11", ("",))[0] == "pd"]
    if spd:
        add("B8", f"Sensitive personal data sits in {plural(len(spd), 'system')}: {join(spd)}.", [f"C11:{s}" for s in spd])
    if pd_:
        add("B8", f"Personal data sits in {plural(len(pd_), 'more system')}: {join(pd_)}.", [f"C11:{s}" for s in pd_])
    vs = [v["Vendor"] for v in vendors if str(v.get("V3 code") or "") == "spd"]
    if vs:
        add("B8", f"Vendors with access to sensitive personal data: {join(vs)}.", [f"V3:{v}" for v in vs])

    pc = [d for d in depts if dcode(d, "U4") and dcode(d, "U4") != "none"]
    if pc:
        add("B8", "Departments that move work data through personal channels: " + join([f"{d['Department']} ({d.get('U4 Personal channels for work data') or ''})" for d in pc]) + ".", [f"U4:{d['Department']}" for d in pc])

    # B9 IT organization
    for k, label in [("ith", "IT headcount"), ("secr", "Security role"), ("outs", "Outsourced IT functions")]:
        if ans(k) not in ("", None):
            add("B9", f"{label}: {ans(k)}.", [f"ENV:{k}"])

    if depts:
        tr = [d for d in depts if dcode(d, "U6") == "12m"]
        add("B9", f"Security training within 12 months: {len(tr)} of {plural(len(depts), 'department')} interviewed.", ["POP:U6"])

    # B10 control coverage snapshot
    names = {r["Field ID"]: r["Field"] for r in prof}
    for fid in COVERAGE_FIELDS:
        c = Counter(f[fid][1] for f in pv.values() if fid in f and f[fid][1] not in ("", "Not applicable"))
        if not c:
            continue
        ca = Counter(f[fid][1] for s, f in pv.items() if crit.get(s) == "A" and fid in f and f[fid][1] not in ("", "Not applicable"))
        add("B10", f"{fid} {names.get(fid, fid)}: {c.get('Good', 0)} Good, {c.get('Weak', 0)} Weak, {c.get('Poor', 0)} Poor, {c.get('Unknown', 0)} Unknown of {sum(c.values())} systems answered"
            + (f" (Criticality A: {ca.get('Good', 0)} Good of {sum(ca.values())})." if ca else "."), [f"POP:{fid}"])

    # B11 assessment basis
    add("B11", f"Assessment date {summary.get('Assessment date') or 'not stated'}. {plural(len(systems), 'system')} listed, {sum(1 for s in systems if s.get('Sampled') == 'Yes')} sampled.", ["Summary:Assessment date", "POP:systems"])
    unk = sum(1 for r in prof if r.get("Rating") == "Unknown")
    add("B11", f"{plural(unk, 'profile answer')} recorded as Unknown.", ["POP:unknown"])
    vul = Counter(o.get("Severity") or "Not rated" for o in obs if o.get("Type") == "Vulnerability")
    if vul:
        add("B11", "Technical observations (vulnerabilities): " + join([f"{v} {k}" for k, v in vul.most_common()]) + ". Details stay in the Technical Observation Register.", [o["ID"] for o in obs if o.get("Type") == "Vulnerability"])

    if depts:
        add("B11", f"{plural(len(depts), 'department')} interviewed in S2C. {plural(len(dcf), 'answer')} conflict with what IT stated.", ["POP:departments"] + [c["ID"] for c in dcf])

    # B12 threat context (client history; the assessor adds sector threats and peer incidents)
    dk = [d["Department"] for d in depts if dcode(d, "U7") == "dk"]
    if dk:
        add("B12", "Staff do not know how to report a suspicious event in " + join(dk) + ".", [f"U7:{n}" for n in dk])
    for h in hist:
        add("B12", f"{h.get('Date') or 'Date not stated'}: {str(h.get('Type') or 'incident').lower()} affecting {h.get('Systems affected') or 'unnamed systems'}. {h.get('Business impact') or ''}".strip().rstrip(".") + ".", [h["ID"]])

    # B13 changes in flight
    for c in chg:
        add("B13", f"{c.get('Project')}: {c.get('Affected systems') or 'no system linked'}, target {c.get('Target date') or 'not set'}, {str(c.get('Status') or 'status not stated').lower()}" + (", within 12 months." if c.get("Within 12 months") == "Yes" else "."), [c["ID"]])

    # B14 prior assessments
    for p in prior:
        add("B14", f"{p.get('Type')} on {p.get('Date') or 'date not stated'} by {p.get('Performed by') or 'not stated'}: {n(p.get('Findings open')) if n(p.get('Findings open')) is not None else 'unknown'} findings open, {n(p.get('Findings closed')) if n(p.get('Findings closed')) is not None else 'unknown'} closed.", [p["ID"]])

    return {
        "schema": "csf-gap-landscape-facts/1", "source": source, "asOf": str(summary.get("Exported") or date.today().isoformat())[:10],
        "blocks": {k: {"name": v, "facts": B[k]} for k, v in BLOCKS.items()},
        "securityTools": [{"id": t["ID"], "capability": t.get("Capability") or "", "product": t.get("Product") or "", "coverage": COVN.get(t.get("Coverage code"), ""),
                           "operatedBy": OPN.get(t.get("Operated by code"), ""), "status": STN.get(t.get("Status code"), ""), "notes": t.get("Notes") or ""} for t in tools],
        "changesInFlight": [{"id": c["ID"], "project": c.get("Project") or "", "systems": [x.strip() for x in str(c.get("Affected systems") or "").split(",") if x.strip()],
                             "targetDate": str(c.get("Target date") or ""), "status": c.get("Status") or "", "within12Months": c.get("Within 12 months") == "Yes"} for c in chg],
        "priorAssessments": [{"id": p["ID"], "type": p.get("Type") or "", "date": str(p.get("Date") or ""), "performedBy": p.get("Performed by") or "",
                              "findingsOpen": n(p.get("Findings open")), "findingsClosed": n(p.get("Findings closed"))} for p in prior],
        "departmentConflicts": [{"id": c["ID"], "department": c.get("Department") or "", "type": c.get("Type code") or "", "detail": c.get("Detail") or "",
                                 "evidence": [e.strip() for e in str(c.get("Evidence") or "").split(",") if e.strip()]} for c in dcf],
        "incidentHistory": [{"id": h["ID"], "date": str(h.get("Date") or ""), "type": h.get("Type") or "", "systems": h.get("Systems affected") or "",
                             "impact": h.get("Business impact") or "", "reported": h.get("Reported to a regulator") or ""} for h in hist],
    }

if __name__ == "__main__":
    out = build(sys.argv[1])
    s = json.dumps(out, indent=1, ensure_ascii=False)
    if len(sys.argv) > 2:
        open(sys.argv[2], "w", encoding="utf-8").write(s + "\n")
        print("wrote", sys.argv[2], sum(len(b["facts"]) for b in out["blocks"].values()), "facts")
    else:
        print(s)
