# -*- coding: utf-8 -*-
"""Acceptance test for Information-Gathering.html before freeze (Information Gathering Spec v2.0 Section 16).
Usage (from _Build-Source): python ig_acceptance.py [../00_Master/Information-Gathering.html]
Runs items 1 to 9, 11 and 12. Item 10 (existing checks) is the rest of run_checks.py.
Needs: playwright, openpyxl, jsonschema. Fixtures: fixtures/cmc-csf-interview-3.json and fixtures/cmc-csf-doc-audit-1.json."""
import asyncio, json, os, sys, re, tempfile, pathlib, base64
import openpyxl
from playwright.async_api import async_playwright

R = pathlib.Path(__file__).resolve().parent
HTML = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else R.parent / "00_Master" / "Information-Gathering.html").resolve()
sys.path.insert(0, str(R.parent / "00_Master" / "Gap-Analysis-Method" / "tools"))
from check_handoff import check as check_handoff
sys.path.insert(0, str(R))
from packs import MANDATES
LOG = R.parent / "00_Master" / "Information-Gathering-Packs-Verification-Log.md"
W = tempfile.mkdtemp()
VIEWS = ["setup", "scope", "plan", "req", "sched", "comms", "sa", "ai", "iv", "sys", "ven", "env", "dept", "tech", "obs", "dapp", "drev", "ev", "fu", "rate", "score", "land", "dash", "ho"]
EXPECT = {"mfg": "S", "hc": "M", "fin": "L"}
res = []


def ok(n, c, info=""):
    res.append(bool(c))
    print(("PASS " if c else "FAIL ") + n + (f"  [{info}]" if info not in ("", None) else ""))


async def dl(pg, sel, name=None):
    async with pg.expect_download() as d:
        await pg.click(sel)
    v = await d.value
    f = os.path.join(W, name or v.suggested_filename)
    await v.save_as(f)
    return f


async def main():
    async with async_playwright() as p:
        b = await p.chromium.launch(executable_path="/opt/pw-browsers/chromium" if os.path.exists("/opt/pw-browsers/chromium") else None)
        ctx = await b.new_context(accept_downloads=True, viewport={"width": 1400, "height": 900})
        await ctx.route("**/fonts.googleapis.com/**", lambda r: r.abort())

        async def page():
            pg = await ctx.new_page()
            pg.errs = []
            pg.on("pageerror", lambda e: pg.errs.append(str(e)))
            await pg.goto(HTML.as_uri())
            await pg.wait_for_timeout(400)
            return pg

        pg = await page()
        # 1, 2, 3, 9: samples end to end, handoff valid, evidence links, no Target
        for k, sc in EXPECT.items():
            await pg.evaluate(f"loadSample('{k}')")
            for v in VIEWS:
                await pg.evaluate(f"setView('{v}')")
            st = await pg.evaluate("[scale(),SUBS.filter(s=>eff(s).st==='none').length,hoCheck().err]")
            ok(f"1 sample {k} runs end to end at scale {sc}", st[0] == sc and not pg.errs, f"scale {st[0]}, errors {pg.errs[:2]}")
            await pg.evaluate("setView('ho')")
            f = await dl(pg, "#b-ho", f"handoff-{k}.json")
            h = json.load(open(f, encoding="utf-8"))
            e = check_handoff(h)
            ok(f"2 sample {k} exports a valid csf-p1-handoff/1 with 106 rated Subcategories", not e and sum(1 for s in h["subcategories"] if s["statedTier"]) == 106, "; ".join(e[:3]))
            nolink = [s["id"] for s in h["subcategories"] if s["evidenceLevel"] in ("D", "O") and not s["evidenceIds"]]
            ok(f"3 sample {k}: no Documented or Demonstrated rating lacks an Evidence ID", not nolink, nolink[:5])
            fx = await dl(pg, "#b-xlsx", f"export-{k}.xlsx")
            wb = openpyxl.load_workbook(fx, read_only=True)
            hdr = [str(c) for ws in wb for row in ws.iter_rows(max_row=1, values_only=True) for c in row if c]
            state = await pg.evaluate("JSON.stringify(S)")
            tg = [x for x in hdr if re.search(r"\btarget\b(?! date)", x, re.I)] + (["state tgt"] if '"tgt"' in state else [])  # "Target date" of a change in flight is a date, not a Target Tier; handoff keys are checked by check_handoff
            ok(f"9 sample {k}: no Target value in state, Excel headers, or handoff", not tg, tg[:3])
        # anonymized handoff
        await pg.evaluate("loadSample('hc');S.setup.anonX=true;setView('ho')")
        h = json.load(open(await dl(pg, "#b-ho", "handoff-anon.json"), encoding="utf-8"))
        txt = json.dumps(h, ensure_ascii=False)
        names = await pg.evaluate("[S.client.name,...S.people.map(p=>p.name)].filter(Boolean)")
        ok("2b aliased handoff hides client and people names and stays valid", not any(n in txt for n in names) and not check_handoff(h) and h["engagement"]["aliased"], [n for n in names if n in txt][:2])

        # 4: merge of two session files with conflicts
        pg = await page()
        A = await pg.evaluate("(()=>{loadSample('hc');return JSON.parse(JSON.stringify(S))})()")
        B = json.loads(json.dumps(A))
        B["subs"]["PR.AA-05"]["tier"] = "3"
        B["subs"]["GV.PO-01"]["notes"] = "Sample: session file note"
        B["people"].append({"name": "Sample Nurse Manager", "role": "SYSOWN", "dept": "Nursing", "date": "2026-10-08"})
        B["subs"]["PR.AA-01"].setdefault("st", []).append({"p": len(B["people"]) - 1, "k": "N", "txt": "Sample: nurses share logins"})
        B["team"].append({"name": "Sample Analyst 2", "ini": "SC"})
        B["ev"].append({"id": "EV-SC-001", "kind": "shot", "title": "Sample: AD admin group", "rdate": "2026-10-07", "recv": "2026-10-07", "subs": ["PR.AA-05"], "docs": [], "tcs": [], "sys": [], "obs": [], "conf": "N"})
        B["sys"].append({"name": "Sample: Nurse call system"})
        fb = os.path.join(W, "session-b.json")
        json.dump(B, open(fb, "w", encoding="utf-8"), ensure_ascii=False)
        await pg.evaluate("setView('setup')")
        await pg.set_input_files("#file2", fb)
        await pg.wait_for_timeout(600)
        cf = await pg.evaluate("MG?MG.cf.map(c=>c.path.join('>')):[]")
        await pg.evaluate("(()=>{const i=MG.cf.findIndex(c=>c.path.join('>')==='subs>PR.AA-05>tier');document.querySelector(`input[name=mg${i}][value=b]`).click();applyMerge();})()")
        r = await pg.evaluate("({t:S.subs['PR.AA-05'].tier,np:S.people.length,ev:S.ev.map(e=>e.id),st:S.subs['PR.AA-01'].st.length,sys:S.sys.length,note:S.subs['GV.PO-01'].notes,open:!!MG})")
        lost = [e["id"] for e in A["ev"] if e["id"] not in r["ev"]]
        ok("4 merge with conflicts resolves without data loss", "subs>PR.AA-05>tier" in cf and r["t"] == "3" and r["np"] == len(A["people"]) + 1 and "EV-SC-001" in r["ev"] and not lost
           and r["st"] == len(A["subs"]["PR.AA-01"].get("st", [])) + 1 and r["sys"] == len(A["sys"]) + 1 and not r["open"], {"conflicts": len(cf), "lost": lost, "people": r["np"]})

        # 5: self-assessment round trip
        await pg.evaluate("loadSample('hc');setView('sa')")
        f = await dl(pg, "#b-saxlsx", "sa.xlsx")
        wb = openpyxl.load_workbook(f)
        q = wb["Questions"]
        rows = {q.cell(i, 1).value: i for i in range(2, q.max_row + 1)}
        q.cell(rows["PR.AA-03"], 4).value = "No"
        q.cell(rows["PR.AA-03"], 5).value = "Client: VPN uses password only"
        sy = wb["Systems"]
        sy.cell(next(i for i in range(2, sy.max_row + 1) if not sy.cell(i, 1).value), 1).value = "Sample: Queue kiosk"
        wb["Documents"].cell(2, 4).value = "Yes"
        fp = os.path.join(W, "sa-filled.xlsx")
        wb.save(fp)
        await pg.set_input_files("#f-sa", fp)
        await pg.wait_for_timeout(800)
        r = await pg.evaluate("(()=>{const z=(S.subs['PR.AA-03'].st||[]).find(z=>z.p===-2);return [z&&z.k,S.sys.some(y=>y.name==='Sample: Queue kiosk'),saDiff().map(d=>d.s)]})()")
        ok("5 self-assessment export, fill, import, difference report", r[0] == "N" and r[1] and "PR.AA-03" in r[2], r)

        # 6: AI prompt imports
        samples = {
            "P-01": {"schema": "csf-ig-pre/1", "facts": {"emp": 1200, "sites": 2, "ot": True}, "flags": {"cii": True}, "industry": "hc", "systems": [{"name": "Sample: Patient portal", "svc": "Web", "source": "https://example.org"}]},
            "P-02": {"schema": "csf-ig-notes/1", "session": "S2", "items": [{"sub": "PR.AA-05", "person": "Person 1", "stance": "P", "statement": "Access review is yearly", "docs": [], "requests": []}], "observations": [{"system": "HIS", "description": "Shared admin account", "severity": "H"}], "followups": []},
            "P-05": {"schema": "csf-ig-land/1", "facts": [{"block": "Backup", "text": "Backups run nightly.", "ev": ["TOOL-09"]}]}}
        did = await pg.evaluate("D.filter(inScope)[0].id")
        cid = await pg.evaluate("dchecks(D.filter(inScope)[0])[0].id")
        samples["P-03"] = {"schema": "csf-ig-docrev/1", "doc": did, "checks": [{"id": cid, "result": "P", "quote": "Approved by the board"}], "status": "Partial", "note": ""}
        await pg.evaluate("S.land.locked=false")
        for pid, j in samples.items():
            n = await pg.evaluate("([p,j])=>{aip=p;const it=aiCheck(j);let n=0;it.forEach(z=>{if(z.a){z.a();n++;}});return n}", [pid, json.dumps(j)])
            ok(f"6 {pid} imports a sample JSON reply", n >= 1, n)
        t4 = await pg.evaluate("(()=>{aip='P-04';return promptText()})()")
        ok("6 P-04 builds a follow-up prompt (no import by design)", "Conflicting answers" in t4 and len(t4) > 200)

        # 7: migration of CMC csf-interview/3 and Audit save files
        iv = json.load(open(R / "fixtures" / "cmc-csf-interview-3.json", encoding="utf-8"))
        au = json.load(open(R / "fixtures" / "cmc-csf-doc-audit-1.json", encoding="utf-8"))
        fe = os.path.join(W, "engagement.json")
        json.dump({"app": "csf-engagement", "v": 1, "interview": iv, "audit": au}, open(fe, "w", encoding="utf-8"), ensure_ascii=False)
        pg = await page()
        await pg.set_input_files("#file", fe)
        await pg.wait_for_timeout(600)
        s = await pg.evaluate("JSON.parse(JSON.stringify(S))")
        tiers_ok = all((s["subs"].get(k) or {}).get("tier", "") == (v or {}).get("tier", "") for k, v in iv["subs"].items())
        docs_ok = all((s["docs"].get(k) or {}).get("status", "") == (v or {}).get("status", "") for k, v in au["docs"].items() if v)
        ok("7 CMC csf-interview/3 and Audit files migrate without loss of ratings or document statuses", tiers_ok and docs_ok and s["schema"] == "csf-ig/4" and '"tgt"' not in json.dumps(s) and not pg.errs,
           f"subs {len(iv['subs'])}, docs {len([1 for v in au['docs'].values() if v])}")

        # 8: password save and open
        await pg.evaluate("loadSample('fin')")
        before = await pg.evaluate("JSON.stringify(Object.assign({},S,{updated:0}))")
        await pg.click("#b-save")
        await pg.fill("#pw1", "Acceptance-Pass-8")
        await pg.fill("#pw2", "Acceptance-Pass-8")
        f = await dl(pg, "#pw-ok", "enc.json")
        enc = json.load(open(f, encoding="utf-8"))
        await pg.evaluate("S=blank();renderAll()")
        await pg.set_input_files("#file", f)
        await pg.wait_for_timeout(300)
        await pg.fill("#pw1", "Acceptance-Pass-8")
        await pg.click("#pw-ok")
        await pg.wait_for_timeout(2000)
        after = await pg.evaluate("JSON.stringify(Object.assign({},S,{updated:0}))")
        ok("8 password save and open round trip", enc.get("app") == "csf-ig-enc" and "Krung" not in open(f, encoding="utf-8").read() and before == after)

        # 11: TH and EN labels, emails, prompts
        miss = await pg.evaluate("(()=>{const out=[];const walk=(a,b,p)=>{for(const k of Object.keys(a)){if(!(k in b)){out.push(p+k);continue;}if(a[k]&&typeof a[k]==='object'&&!Array.isArray(a[k]))walk(a[k],b[k],p+k+'.');}};walk(T.EN,T.TH,'TH:');walk(T.TH,T.EN,'EN:');"
                                 "COMMS.forEach(c=>['subj','body','n','when'].forEach(k=>{if(!c[k][0]||!c[k][1])out.push(c.id+'.'+k)}));PROMPTS.forEach(p=>['n','d'].forEach(k=>{if(!p[k][0]||!p[k][1])out.push(p.id+'.'+k)}));return out})()")
        ok("11 all labels, emails, and prompt names exist in TH and EN (prompt bodies are English by design)", not miss, miss[:5])
        await b.close()

    # 12: legal references verified or marked [VERIFY]
    log = open(LOG, encoding="utf-8").read()
    bad = []
    for m in MANDATES:
        row = next((l for l in log.splitlines() if l.startswith(f"| {m['id']} |")), "")
        want = "Instrument checked" if m["ver"] == "verified" else "Verify clause"
        if m["ver"] not in ("verified", "verify") or want not in row:
            bad.append(m["id"])
    ok("12 every pack legal reference is Verified or marked Verify in the log", not bad, bad)
    print(f"RESULT {'PASS' if all(res) else 'FAIL'} {sum(res)}/{len(res)}")
    sys.exit(0 if all(res) else 1)


asyncio.run(main())
