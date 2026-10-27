# -*- coding: utf-8 -*-
import os, re, glob, sys
from docx import Document
from meta import DOC, DOCS
from build import load_content

OUT = "out/NIST-CSF2-Document-Framework"
issues = []
def text_of(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                parts.append(c.text)
    for s in d.sections:
        parts += [p.text for p in s.header.paragraphs]
    return "\n".join(parts)

BANNED = r"\b(estate|posture|uplift|leverag\w*|artefact\w*)\b"
BRIT = r"\b(organisation\w*|behaviour\w*|colour\w*|prioritis\w*|authoris\w*|analys(e|ed|es|ing)\b|centre\w*|programme\w*|defence|catalogue\w*|minimis\w*|recognis\w*|utilis\w*|licence)\b"
IDPAT = r"\b(GV|ID|PR|DE|RS|RC)-(CHR|STR|STM|POL|STD|PRC|PLN|PLB|GDL|FRM|REG|REC|RPT|TPL)-\d{3}\b"
EMOJI = re.compile("[\U0001F300-\U0001FAFF☀-➿]")
allowed_cmc_ph = {"RS-TPL-002", "RC-TPL-001", "GV-PRC-002", "GV-PRC-001"}

C = load_content()
# content structure checks
for d in DOCS:
    c = C[d["id"]]
    if not c.get("bp"): issues.append((d["id"], "no best practice"))
    if not c.get("mistakes"): issues.append((d["id"], "no mistakes"))
    if d["type"] in ("FRM",) and not c.get("fields"): issues.append((d["id"], "form without fields"))
    if d["type"] in ("REG",) and not c.get("fields"): issues.append((d["id"], "register without fields"))
    if c.get("fields") and d["type"] in ("REG", "REC"):
        for r in c.get("rows", []):
            if len(r) != len(c["fields"]): issues.append((d["id"], f"row len {len(r)} != {len(c['fields'])}"))
        if not c.get("rows"): issues.append((d["id"], "no CMC rows"))
    if c.get("fields") and d["type"] == "FRM":
        if len(c.get("cmc_vals", [])) != len(c["fields"]): issues.append((d["id"], "cmc_vals mismatch"))
    for s in c.get("sections", []):
        if len(s) < 6: issues.append((d["id"], f"section short {s[0]}"))

files = glob.glob(f"{OUT}/**/*.docx", recursive=True)
counts = {}
for f in files:
    setname = f.split("/")[2]
    did = os.path.basename(f)[:10]
    t = text_of(f)
    counts[setname] = counts.get(setname, 0) + 1
    if "—" in t or "–" in t: issues.append((f, "dash found"))
    if EMOJI.search(t): issues.append((f, "emoji"))
    for m in re.finditer(IDPAT, t):
        if m.group(0) not in DOC: issues.append((f, "bad ref " + m.group(0)))
    if setname != "01_Template_TH":
        for m in re.finditer(BANNED, t, re.I): issues.append((f, "banned: " + m.group(0)))
        for m in re.finditer(BRIT, t, re.I): issues.append((f, "brit: " + m.group(0)))
        if re.search("[฀-๿]", t): issues.append((f, "thai chars in EN"))
    else:
        th = len(re.findall("[฀-๿]", t)); lat = len(re.findall("[A-Za-z]", t))
        if th < lat * 0.6: issues.append((f, f"low thai ratio th={th} lat={lat}"))
        for m in re.finditer(BANNED, t, re.I): issues.append((f, "banned: " + m.group(0)))
    if setname == "03_CMC_EN":
        ph = set(re.findall(r"\[[^\]]+\]", t)) - {"[VERIFY]"}
        ph = {p for p in ph if not p.startswith("[VERIFY")}
        if ph and did not in allowed_cmc_ph: issues.append((f, "placeholders: " + ", ".join(sorted(ph))[:300]))
print(counts)
print(len(issues), "issues")
for i in issues[:300]:
    print(i)
