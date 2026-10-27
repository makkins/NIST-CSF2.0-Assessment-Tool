# -*- coding: utf-8 -*-
"""Render all documents into three sets: Template TH, Template EN, CMC EN."""
import importlib, glob, os, re, sys, datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from dictionary import PKG_TH
from meta import (DOCS, DOC, ROLES, CMC_REPLACE, FUNCTIONS, CATEGORIES, SUBCATS, TYPES, TIERS, PPT,
                  LEVEL_TH, TRIGGER_TH, cmc_level)

OUT = os.path.join(os.path.dirname(__file__), "out", "NIST-CSF2-Document-Framework")
NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x59, 0x59, 0x59)
REQ_FILL = "EAF0F7"
HEAD_FILL = "1F3864"
SAMPLE_FILL = "FFFFFF"

# ---------------------------------------------------------------- content
def load_content():
    C = {}
    for f in sorted(glob.glob(os.path.join(os.path.dirname(__file__), "content", "c_*.py"))):
        name = os.path.splitext(os.path.basename(f))[0]
        m = importlib.import_module("content." + name)
        for k, v in m.C.items():
            assert k in DOC, (name, k)
            assert k not in C, ("dup", k)
            C[k] = v
    return C

def slug(s):
    s = re.sub(r"[()]", "", s)
    s = re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-")
    return s

def cmc_text(s):
    if s is None:
        return None
    if isinstance(s, str):
        for a, b in CMC_REPLACE:
            s = s.replace(a, b)
        # role placeholders like [CISO]
        for k, v in ROLES.items():
            s = s.replace("[" + k + "]", v[2])
        return s
    if isinstance(s, list):
        return [cmc_text(x) for x in s]
    if isinstance(s, tuple):
        return tuple(cmc_text(x) for x in s)
    if isinstance(s, dict):
        return {k: cmc_text(v) for k, v in s.items()}
    return s

def role_text(s, lang, mode):
    """Replace [ROLE] placeholders for template sets."""
    if s is None:
        return s
    if isinstance(s, str):
        for k, v in ROLES.items():
            if mode == "cmc":
                rep = v[2]
            else:
                rep = v[1] if lang == "TH" else v[0]
            s = s.replace("[" + k + "]", rep)
        return s
    if isinstance(s, list):
        return [role_text(x, lang, mode) for x in s]
    if isinstance(s, tuple):
        return tuple(role_text(x, lang, mode) for x in s)
    if isinstance(s, dict):
        return {k: role_text(v, lang, mode) for k, v in s.items()}
    return s

# ---------------------------------------------------------------- docx helpers
class R:
    def __init__(self, lang):
        self.lang = lang
        self.font = "TH Sarabun New" if lang == "TH" else "Calibri"
        self.size = 15 if lang == "TH" else 10.5
        self.small = 13 if lang == "TH" else 9
        self.tiny = 12 if lang == "TH" else 8

def set_run_font(run, r, size=None, bold=None, italic=None, color=None):
    run.font.name = r.font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), r.font)
    sz = size or r.size
    run.font.size = Pt(sz)
    szcs = rpr.find(qn("w:szCs"))
    if szcs is None:
        szcs = OxmlElement("w:szCs"); rpr.append(szcs)
    szcs.set(qn("w:val"), str(int(sz * 2)))
    if bold is not None:
        run.font.bold = bold
        if bold:
            b = OxmlElement("w:bCs"); rpr.append(b)
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color

def shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def para(doc_or_cell, text, r, size=None, bold=False, italic=False, color=None, style=None, align=None, space_after=4):
    p = doc_or_cell.add_paragraph(style=style) if style else doc_or_cell.add_paragraph()
    if text:
        add_rich(p, text, r, size=size, bold=bold, italic=italic, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align:
        p.alignment = align
    return p

def add_rich(p, text, r, size=None, bold=False, italic=False, color=None):
    """Support **bold** segments."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2]); set_run_font(run, r, size, True, italic, color)
        else:
            run = p.add_run(part); set_run_font(run, r, size, bold, italic, color)

def heading(doc, text, level, r):
    sizes = {0: (r.size + 9), 1: (r.size + 4), 2: (r.size + 1.5), 3: r.size}
    p = doc.add_heading(level=min(level, 3) if level else 1)
    run = p.add_run(text)
    set_run_font(run, r, sizes[level], True, False, NAVY)
    p.paragraph_format.space_before = Pt(10 if level <= 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    return p

def bullet(doc_or_cell, text, r, size=None, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    p = doc_or_cell.add_paragraph(style=style)
    add_rich(p, text, r, size=size)
    p.paragraph_format.space_after = Pt(2)
    return p

def set_col_widths(table, widths_cm):
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for row in table.rows:
        for i, w in enumerate(widths_cm):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)
    grid = tbl.tblGrid
    cols = grid.findall(qn("w:gridCol"))
    for i, w in enumerate(widths_cm):
        if i < len(cols):
            cols[i].set(qn("w:w"), str(int(w * 567)))

def table(doc, headers, rows, r, widths=None, size=None, header_fill=HEAD_FILL):
    size = size or r.small
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        add_rich(c.paragraphs[0], str(h), r, size=size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(c, header_fill)
    # repeat header row
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i in range(len(headers)):
            v = row[i] if i < len(row) else ""
            cells[i].text = ""
            txt = "" if v is None else str(v)
            lines = txt.split("\n")
            add_rich(cells[i].paragraphs[0], lines[0], r, size=size)
            for ln in lines[1:]:
                pp = cells[i].add_paragraph(); add_rich(pp, ln, r, size=size)
    if not widths:
        widths = [16.5 / len(headers)] * len(headers)
    set_col_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def kv_table(doc, rows, r, widths=(5.0, 11.5)):
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in rows:
        cells = t.add_row().cells
        add_rich(cells[0].paragraphs[0], k, r, size=r.small, bold=True)
        shade(cells[0], REQ_FILL)
        lines = str(v).split("\n")
        add_rich(cells[1].paragraphs[0], lines[0], r, size=r.small)
        for ln in lines[1:]:
            pp = cells[1].add_paragraph(); add_rich(pp, ln, r, size=r.small)
    set_col_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def box(doc, label, text, r, fill=REQ_FILL):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    shade(c, fill)
    p = c.paragraphs[0]
    add_rich(p, label + " ", r, size=r.small, bold=True, color=NAVY)
    add_rich(p, text, r, size=r.small)
    set_col_widths(t, [16.5])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def render_body(doc, body, r):
    if body is None or body == "":
        return
    if isinstance(body, str):
        for ln in body.split("\n"):
            if ln.strip():
                para(doc, ln, r)
    elif isinstance(body, list):
        for it in body:
            if isinstance(it, (dict, tuple)):
                render_body(doc, it, r)
            else:
                bullet(doc, it, r)
    elif isinstance(body, dict):
        if "t" in body:
            table(doc, body["t"], body["r"], r, widths=body.get("w"), size=r.tiny if len(body["t"]) > 5 else None)
        elif "n" in body:
            for it in body["n"]:
                bullet(doc, it, r, numbered=True)
        elif "seq" in body:
            for b in body["seq"]:
                render_body(doc, b, r)
        elif "p" in body:
            render_body(doc, body["p"], r)
    elif isinstance(body, tuple):
        for b in body:
            render_body(doc, b, r)

def add_page_number(par, r, prefix):
    add_rich(par, prefix, r, size=r.tiny, color=GRAY)
    for code in ("PAGE",):
        run = par.add_run(); set_run_font(run, r, r.tiny, color=GRAY)
        f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
        it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = code
        f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
        run._r.append(f1); run._r.append(it); run._r.append(f2)

# ---------------------------------------------------------------- labels
L = {
 "EN": dict(doc_control="Document Control", profile="Document Profile", rev="Revision History",
            partA="Part A: Requirement Fields", partB="Part B: Best Practice", mistakes="Common Mistakes to Avoid",
            annex="Annex: Traceability and Related Documents", req="Requirement:", sample="Sample content",
            id="Document ID", title="Title", version="Version", status="Status", cls="Classification",
            owner="Document Owner", approver="Approver", eff="Effective Date", cycle="Review Cycle", next="Next Review",
            tier="Tier", type="Document Type", fn="CSF Function", cat="CSF Category", level="Level", trigger="Trigger",
            basis="Basis", ppt="ISO/IEC 27002 theme and tags", pkg="Package", cmclevel="Level for CMC",
            ver="Version", date="Date", change="Description of Change", author="Prepared By", appr="Approved By",
            csf="NIST CSF 2.0 Subcategories", iso="ISO/IEC 27001:2022", nist="NIST SP 800-53 Rev. 5",
            refs="Legal, Regulatory, and Other References", parents="Parent Documents", children="Child Documents",
            none="None", internal="Internal", howto="How to Use This Template", about="About This Sample",
            fields="Form Fields", field="Field", desc="Description", entry="Entry", example="Example",
            cols="Column Definitions", column="Column", entries="Sample Entries", changelog="Register Change Log",
            approval="Approval and Sign-off", role="Role", name="Name", sig="Signature", records="Records Produced",
            purpose="Purpose", scope="Scope", compliance="Compliance and Exceptions", page="Page ",
            draft="Template (not approved)", approved="Approved"),
 "TH": dict(doc_control="การควบคุมเอกสาร", profile="ข้อมูลประจำเอกสาร", rev="ประวัติการแก้ไข",
            partA="ส่วน A: ข้อกำหนดขั้นต่ำ (Requirement Fields)", partB="ส่วน B: แนวปฏิบัติที่ดี (Best Practice)",
            mistakes="ข้อผิดพลาดที่พบบ่อย", annex="ภาคผนวก: การอ้างอิงและเอกสารที่เกี่ยวข้อง", req="ข้อกำหนด:",
            sample="ตัวอย่างเนื้อหา", id="รหัสเอกสาร", title="ชื่อเอกสาร", version="ฉบับที่", status="สถานะ",
            cls="ระดับชั้นความลับ", owner="เจ้าของเอกสาร", approver="ผู้อนุมัติ", eff="วันที่มีผลบังคับใช้",
            cycle="รอบการทบทวน", next="วันที่ทบทวนครั้งถัดไป", tier="ชั้นเอกสาร", type="ประเภทเอกสาร",
            fn="CSF Function", cat="CSF Category", level="ระดับความจำเป็น (Level)", trigger="เงื่อนไขที่ทำให้บังคับ (Trigger)",
            basis="แหล่งอ้างอิง (Basis)", ppt="ISO/IEC 27002 theme และ Tag", pkg="ชุดเอกสาร (Package)", cmclevel="",
            ver="ฉบับที่", date="วันที่", change="รายละเอียดการแก้ไข", author="ผู้จัดทำ", appr="ผู้อนุมัติ",
            csf="NIST CSF 2.0 Subcategories", iso="ISO/IEC 27001:2022", nist="NIST SP 800-53 Rev. 5",
            refs="กฎหมาย ระเบียบ และแหล่งอ้างอิงอื่น", parents="เอกสารแม่ (Parent)", children="เอกสารลูก (Child)",
            none="ไม่มี", internal="ใช้ภายใน (Internal)", howto="วิธีใช้แม่แบบนี้", about="",
            fields="รายการข้อมูลในแบบฟอร์ม", field="หัวข้อ", desc="คำอธิบาย", entry="ข้อมูลที่กรอก", example="ตัวอย่าง",
            cols="คำอธิบายคอลัมน์", column="คอลัมน์", entries="ตัวอย่างรายการ", changelog="บันทึกการเปลี่ยนแปลงทะเบียน",
            approval="การอนุมัติและลงนาม", role="บทบาท", name="ชื่อ", sig="ลายมือชื่อ", records="บันทึกที่เกิดจากกระบวนการนี้",
            purpose="วัตถุประสงค์", scope="ขอบเขต", compliance="การปฏิบัติตามและการขอยกเว้น", page="หน้า ",
            draft="แม่แบบ (ยังไม่ได้อนุมัติ)", approved="อนุมัติแล้ว"),
}


def type_name(d, lang):
    return TYPES[d["type"]][1] if lang == "TH" else TYPES[d["type"]][0]

def doc_name(did, lang):
    d = DOC[did]
    return d["th"] if lang == "TH" else d["en"]

# ---------------------------------------------------------------- auto sections
def auto_scope(d, lang):
    t = d["type"]
    tn = type_name(d, lang).lower() if lang == "EN" else type_name(d, lang)
    if lang == "EN":
        if t in ("POL", "STD"):
            return f"This {tn} applies to all employees, contractors, temporary staff, and third parties who access [Organization Name] information or information systems, at all sites, including cloud services operated on behalf of [Organization Name]."
        if t in ("CHR", "STR", "STM"):
            return f"This {tn} covers [Organization Name] as a whole, including all business units, sites, and outsourced services."
        if t in ("PRC", "PLN", "PLB", "GDL"):
            return f"This {tn} applies to all units and personnel who perform or support the activities described below, including third parties acting on behalf of [Organization Name]."
        par = ", ".join(d["parents"])
        return f"This {tn} is used by the units and roles named in its parent document ({par}) and covers the activities that document defines."
    else:
        if t in ("POL", "STD"):
            return f"{tn}นี้ใช้บังคับกับพนักงาน ผู้รับจ้าง พนักงานชั่วคราว และบุคคลภายนอกทุกรายที่เข้าถึงข้อมูลหรือระบบสารสนเทศของ [ชื่อองค์กร] ในทุกสถานที่ รวมถึงบริการคลาวด์ที่ดำเนินการในนามของ [ชื่อองค์กร]"
        if t in ("CHR", "STR", "STM"):
            return f"{tn}นี้ครอบคลุม [ชื่อองค์กร] ทั้งองค์กร รวมถึงทุกหน่วยธุรกิจ ทุกสถานที่ และบริการที่จ้างภายนอก"
        if t in ("PRC", "PLN", "PLB", "GDL"):
            return f"{tn}นี้ใช้กับทุกหน่วยงานและบุคลากรที่ดำเนินการหรือสนับสนุนกิจกรรมที่ระบุในเอกสารนี้ รวมถึงบุคคลภายนอกที่ปฏิบัติงานในนามของ [ชื่อองค์กร]"
        par = ", ".join(d["parents"])
        return f"{tn}นี้ใช้โดยหน่วยงานและบทบาทที่ระบุในเอกสารแม่ ({par}) และครอบคลุมกิจกรรมตามที่เอกสารแม่กำหนด"

AUTO_REQ = {
 "purpose": ("State why the document exists and the outcome it supports.", "ระบุเหตุผลของการมีเอกสารนี้และผลลัพธ์ที่เอกสารสนับสนุน"),
 "scope": ("Define the units, people, systems, locations, and third parties covered, and state any exclusions.", "ระบุหน่วยงาน บุคลากร ระบบ สถานที่ และบุคคลภายนอกที่อยู่ในขอบเขต พร้อมระบุข้อยกเว้น (ถ้ามี)"),
 "compliance": ("State that compliance is mandatory, how compliance is verified, how exceptions are requested, and the consequence of violation.", "ระบุว่าการปฏิบัติตามเป็นข้อบังคับ วิธีการตรวจสอบการปฏิบัติตาม ช่องทางการขอยกเว้น และผลของการฝ่าฝืน"),
 "records": ("List the forms, records, and registers this process produces and where they are kept.", "ระบุแบบฟอร์ม บันทึก และทะเบียนที่เกิดจากกระบวนการนี้ และสถานที่จัดเก็บ"),
 "fields": ("List every field the requester or approver must complete, with a clear description, so the form is complete and auditable.", "ระบุทุกหัวข้อที่ผู้ขอหรือผู้อนุมัติต้องกรอก พร้อมคำอธิบายที่ชัดเจน เพื่อให้แบบฟอร์มครบถ้วนและตรวจสอบได้"),
 "approval": ("Record who reviewed and approved the form, with name, date, and signature or electronic approval.", "บันทึกผู้ตรวจสอบและผู้อนุมัติ พร้อมชื่อ วันที่ และลายมือชื่อหรือการอนุมัติทางอิเล็กทรอนิกส์"),
 "cols": ("Define every column, its meaning, and allowed values so entries stay consistent and comparable.", "กำหนดทุกคอลัมน์ ความหมาย และค่าที่อนุญาต เพื่อให้รายการมีความสม่ำเสมอและเปรียบเทียบได้"),
 "entries": ("Maintain current entries. Each entry must have an owner and a last-updated date.", "บันทึกรายการที่เป็นปัจจุบัน ทุกรายการต้องมีผู้รับผิดชอบและวันที่ปรับปรุงล่าสุด"),
 "changelog": ("Track changes to the register itself (additions, closures, and periodic reviews) so reviewers can see how it evolved.", "บันทึกการเปลี่ยนแปลงของทะเบียน (การเพิ่ม การปิด และการทบทวนตามรอบ) เพื่อให้ผู้ตรวจสอบเห็นพัฒนาการของทะเบียน"),
}

def compliance_text(d, lang):
    tn = type_name(d, lang)
    if lang == "EN":
        return [f"Compliance with this {tn.lower()} is mandatory for everyone in scope.",
                "The [CISO] verifies compliance through periodic reviews, technical checks, and internal audit (GV-PRC-007).",
                "Requests for deviation follow the Policy Exception Management Procedure (GV-PRC-003). No exception is valid without written approval and an expiry date.",
                "Violations may result in disciplinary action under the Human Resources Security Standard (GV-STD-001) and applicable law."]
    return [f"บุคคลทุกรายที่อยู่ในขอบเขตต้องปฏิบัติตาม{tn}นี้",
            "[CISO] ตรวจสอบการปฏิบัติตามผ่านการทบทวนตามรอบ การตรวจสอบทางเทคนิค และการตรวจสอบภายใน (GV-PRC-007)",
            "การขอยกเว้นให้ดำเนินการตามระเบียบปฏิบัติการขอยกเว้นการปฏิบัติตามนโยบาย (GV-PRC-003) ข้อยกเว้นจะมีผลเมื่อได้รับอนุมัติเป็นลายลักษณ์อักษรและมีวันหมดอายุเท่านั้น",
            "การฝ่าฝืนอาจนำไปสู่การลงโทษทางวินัยตามมาตรฐานความมั่นคงปลอดภัยด้านทรัพยากรบุคคล (GV-STD-001) และกฎหมายที่เกี่ยวข้อง"]

def records_text(d, lang):
    kids = [c for c in d["children"] if DOC[c]["tier"] == "T4"]
    if kids:
        return [f"{c}: {doc_name(c, lang)}" for c in kids] + (
            ["Records are retained according to the Data Retention and Disposal Standard (PR-STD-007)."] if lang == "EN"
            else ["จัดเก็บบันทึกตามมาตรฐานการเก็บรักษาและการทำลายข้อมูล (PR-STD-007)"])
    return ("Records produced under this document are retained according to the Data Retention and Disposal Standard (PR-STD-007)."
            if lang == "EN" else "บันทึกที่เกิดจากเอกสารนี้ให้จัดเก็บตามมาตรฐานการเก็บรักษาและการทำลายข้อมูล (PR-STD-007)")

# ---------------------------------------------------------------- CMC revision history
CMC_DEFAULT_REV = {
 "GV": "Updated after the 2026 Management Review (MR-2026-01) to reflect the revised budget and ISSC decisions.",
 "ID": "Updated after the November 2025 internal audit (IA-2025-04) to close findings on asset and risk records.",
 "PR": "Updated after the November 2025 internal audit (IA-2025-04) and the 2026 control gap review.",
 "DE": "Updated after onboarding the managed security service provider (MSSP) for 24x7 monitoring in March 2026.",
 "RS": "Updated after the February 2026 ransomware tabletop exercise (EX-2026-01) and the June 2026 email compromise attempt (INC-2026-014).",
 "RC": "Updated after the February 2026 ransomware tabletop exercise (EX-2026-01) and the May 2026 DR test.",
}
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]

def cmc_dates(idx, d):
    # v1.0 between March and September 2025, v1.1 between March and September 2026
    m0 = 3 + (idx * 5) % 7
    day0 = 1 + (idx * 7) % 27
    m1 = 3 + (idx * 3) % 7
    day1 = 1 + (idx * 11) % 27
    return (f"{day0} {MONTHS[m0-1]} 2025", f"{day1} {MONTHS[m1-1]} 2026", f"{day1} {MONTHS[m1-1]} 2027")

# ---------------------------------------------------------------- main render
def render(d, c, lang, mode, idx, outdir):
    r = R(lang)
    lb = L[lang]
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    st = doc.styles["Normal"]
    st.font.name = r.font; st.font.size = Pt(r.size)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), r.font)

    T = lambda s: transform(s, lang, mode)
    org = {"tpl": ("[Organization Name]" if lang == "EN" else "[ชื่อองค์กร]"), "cmc": "Chaophraya Medical Center (CMC)"}[mode]
    title = doc_name(d["id"], lang)

    # header / footer
    hp = sec.header.paragraphs[0]
    add_rich(hp, f"{org}  |  {d['id']}  |  {title}", r, size=r.tiny, color=GRAY)
    fp = sec.footer.paragraphs[0]
    add_page_number(fp, r, f"{lb['cls']}: {lb['internal']}    {lb['page']}")

    # title block
    para(doc, org, r, size=r.size + 1, bold=True, color=GRAY, space_after=2)
    heading(doc, title, 0, r)
    para(doc, f"{d['id']}  |  {type_name(d, lang)}  |  {CATEGORIES[d['cat']][1 if lang=='TH' else 0]} ({d['cat']})", r, size=r.small, color=GRAY)

    if mode == "tpl":
        if lang == "EN":
            box(doc, lb["howto"] + ":", "Replace text in square brackets with your organization's details. Keep every Part A section, because each one is a minimum requirement. Use Part B to raise maturity. Delete the blue requirement boxes and this note before you issue the document. This template is a sample and is not legal advice; confirm items marked [VERIFY] with current Thai regulations before use.", r)
        else:
            box(doc, lb["howto"] + ":", "แทนที่ข้อความในวงเล็บเหลี่ยมด้วยข้อมูลขององค์กร คงทุกหัวข้อในส่วน A ไว้เพราะเป็นข้อกำหนดขั้นต่ำ ใช้ส่วน B เพื่อยกระดับความพร้อม ลบกล่องข้อกำหนดสีฟ้าและหมายเหตุนี้ก่อนประกาศใช้ แม่แบบนี้เป็นตัวอย่างและไม่ใช่คำปรึกษาทางกฎหมาย ให้ตรวจสอบรายการที่ระบุ [VERIFY] กับกฎระเบียบไทยฉบับปัจจุบันก่อนใช้งาน", r)
    else:
        box(doc, lb["about"] + ":", "This sample shows how Chaophraya Medical Center (CMC), a fictional private hospital group in Thailand, applies the template. CMC operates two hospitals with about 300 beds, 1,200 staff, and an IT team of 25. CMC is designated Critical Information Infrastructure (CII) in the public health sector and processes sensitive health data under the PDPA. All names, systems, vendors, and figures are fictional.", r)

    # document control
    v10, v11, nxt = cmc_dates(idx, d)
    heading(doc, "1. " + lb["doc_control"], 1, r)
    own = ROLES[d["owner"]][2 if mode == "cmc" else (1 if lang == "TH" else 0)]
    apr = ROLES[d["approver"]][2 if mode == "cmc" else (1 if lang == "TH" else 0)]
    review = d["review"]
    if lang == "TH":
        review = {"Annual": "ทุกปี", "Quarterly": "ทุกไตรมาส", "Semi-annual": "ทุก 6 เดือน", "Continuous": "ต่อเนื่อง", "Each meeting": "ทุกครั้งที่ประชุม",
                  "Each test": "ทุกครั้งที่ทดสอบ", "Each exercise": "ทุกครั้งที่ฝึกซ้อม", "Each review cycle": "ทุกรอบการทบทวน"}.get(review, review)
        review = (review.replace("Continuous; reconcile quarterly", "ต่อเนื่อง และกระทบยอดทุกไตรมาส")
                  .replace("Continuous; review monthly", "ต่อเนื่อง และทบทวนทุกเดือน")
                  .replace("Continuous; report monthly", "ต่อเนื่อง และรายงานทุกเดือน")
                  .replace("Annual and on legal change", "ทุกปี และเมื่อกฎหมายเปลี่ยนแปลง")
                  .replace("Annual and on change", "ทุกปี และเมื่อมีการเปลี่ยนแปลง")
                  .replace("Annual and on major change", "ทุกปี และเมื่อมีการเปลี่ยนแปลงสำคัญ")
                  .replace("Annual and after each test", "ทุกปี และหลังการทดสอบทุกครั้ง"))
    if mode == "cmc":
        dc = [(lb["id"], d["id"]), (lb["title"], title), (lb["version"], "1.1"), (lb["status"], lb["approved"]),
              (lb["cls"], lb["internal"]), (lb["owner"], own), (lb["approver"], apr), (lb["eff"], v11),
              (lb["cycle"], review), (lb["next"], nxt)]
    else:
        dc = [(lb["id"], d["id"]), (lb["title"], title), (lb["version"], "1.0"), (lb["status"], lb["draft"]),
              (lb["cls"], lb["internal"]), (lb["owner"], own), (lb["approver"], apr),
              (lb["eff"], "[DD Month YYYY]" if lang == "EN" else "[วว เดือน ปปปป]"), (lb["cycle"], review),
              (lb["next"], "[DD Month YYYY]" if lang == "EN" else "[วว เดือน ปปปป]")]
    kv_table(doc, dc, r)

    # profile
    heading(doc, "2. " + lb["profile"], 1, r)
    fn = FUNCTIONS[d["fn"]]
    lvl = d["level"]; trig = d["trigger"]
    if lang == "TH":
        prof = [(lb["tier"], TIERS[d["tier"]][1]), (lb["type"], type_name(d, lang)),
                (lb["fn"], f"{fn[1]} ({fn[0]}, {d['fn']})"), (lb["cat"], f"{CATEGORIES[d['cat']][1]} ({d['cat']})"),
                (lb["level"], LEVEL_TH[lvl]), (lb["trigger"], ", ".join(TRIGGER_TH.get(t, t) for t in trig)),
                (lb["basis"], th_basis(d["basis"])), (lb["ppt"], theme_text(d, "TH")), (lb["pkg"], PKG_TH[pkg_of(d)])]
    else:
        prof = [(lb["tier"], TIERS[d["tier"]][0]), (lb["type"], type_name(d, lang)),
                (lb["fn"], f"{fn[0]} ({d['fn']})"), (lb["cat"], f"{CATEGORIES[d['cat']][0]} ({d['cat']})"),
                (lb["level"], lvl), (lb["trigger"], ", ".join(trig)), (lb["basis"], d["basis"]),
                (lb["ppt"], theme_text(d, "EN")), (lb["pkg"], pkg_of(d))]
        if mode == "cmc":
            prof.append((lb["cmclevel"], cmc_level(d) + " (CMC triggers: PDPA, CII, Computer Crime Act service provider; ISO/IEC 27001 certification targeted for 2027)"))
    kv_table(doc, prof, r)

    # revision history
    heading(doc, "3. " + lb["rev"], 1, r)
    hdr = [lb["ver"], lb["date"], lb["change"], lb["author"], lb["appr"]]
    if mode == "cmc":
        rev11 = c.get("rev") or CMC_DEFAULT_REV[d["fn"]]
        rows = [["1.0", v10, "Initial release under the CMC security program launched in January 2025.", own, apr],
                ["1.1", v11, rev11, own, apr]]
    else:
        if lang == "EN":
            rows = [["1.0", "[DD Month YYYY]", "Initial release", "[Name, Role]", "[Name, Role]"]]
        else:
            rows = [["1.0", "[วว เดือน ปปปป]", "ประกาศใช้ครั้งแรก", "[ชื่อ ตำแหน่ง]", "[ชื่อ ตำแหน่ง]"]]
    table(doc, hdr, rows, r, widths=[1.6, 2.8, 6.6, 2.8, 2.7])

    # PART A
    heading(doc, "4. " + lb["partA"], 1, r)
    if mode == "cmc" and c.get("cmc_note"):
        box(doc, "CMC context:", c["cmc_note"], r, fill="F2F2F2")
    n = 0
    def sect(h, req, body, kind=None):
        nonlocal n
        n += 1
        heading(doc, f"4.{n} {h}", 2, r)
        box(doc, lb["req"], req, r)
        if kind is None:
            p = para(doc, lb["sample"], r, size=r.small, italic=True, color=GRAY, space_after=2)
        render_body(doc, body, r)

    li = 1 if lang == "TH" else 0
    purpose = c["purpose"][li] if mode == "tpl" else (c.get("purpose_cmc") or c["purpose"][0])
    sect(lb["purpose"], AUTO_REQ["purpose"][li], T(purpose))
    scope = c.get("scope")
    scope_txt = (scope[li] if scope else auto_scope(d, lang)) if mode == "tpl" else (c.get("scope_cmc") or (scope[0] if scope else auto_scope(d, "EN")))
    sect(lb["scope"], AUTO_REQ["scope"][li], T(scope_txt))

    for s in c.get("sections", []):
        h = s[1] if lang == "TH" else s[0]
        req = s[3] if lang == "TH" else s[2]
        if mode == "tpl":
            body = s[5] if lang == "TH" else s[4]
        else:
            body = s[6] if len(s) > 6 and s[6] is not None else s[4]
        sect(h, T(req), T(body))

    # forms
    if d["type"] == "FRM" and c.get("fields"):
        n += 1
        heading(doc, f"4.{n} {lb['fields']}", 2, r)
        box(doc, lb["req"], AUTO_REQ["fields"][li], r)
        rows = []
        vals = c.get("cmc_vals") or []
        for i, f in enumerate(c["fields"]):
            fname = f[1] if lang == "TH" else f[0]
            fdesc = f[3] if lang == "TH" else f[2]
            if mode == "cmc":
                entry = vals[i] if i < len(vals) else ""
            else:
                ex = ex_val(f, li)
                entry = (("For example: " if lang == "EN" else "ตัวอย่าง: ") + ex) if ex else ""
            rows.append([str(i + 1), T(fname), T(fdesc), T(entry)])
        table(doc, ["No." if lang == "EN" else "ลำดับ", lb["field"], lb["desc"], lb["entry"]], rows, r, widths=[1.1, 4.0, 6.4, 5.0])
        n += 1
        heading(doc, f"4.{n} {lb['approval']}", 2, r)
        box(doc, lb["req"], AUTO_REQ["approval"][li], r)
        appr_roles = c.get("approvals") or ["Requester", "Line Manager", d["approver"]]
        arows = []
        for ar in appr_roles:
            if ar in ROLES:
                nm = ROLES[ar][2 if mode == "cmc" else li]
            else:
                nm = {"Requester": ("Requester", "ผู้ขอ"), "Line Manager": ("Line Manager", "ผู้บังคับบัญชา"),
                      "Reviewer": ("Reviewer", "ผู้ตรวจสอบ")}.get(ar, (ar, ar))[li]
            arows.append([nm, "", "", ""])
        table(doc, [lb["role"], lb["name"], lb["sig"], lb["date"]], arows, r, widths=[5.0, 4.5, 4.0, 3.0])

    # registers / records / reports with fields
    if d["type"] in ("REG", "REC") and c.get("fields"):
        n += 1
        heading(doc, f"4.{n} {lb['cols']}", 2, r)
        box(doc, lb["req"], AUTO_REQ["cols"][li], r)
        rows = []
        for f in c["fields"]:
            ex = ex_val(f, li)
            rows.append([T(f[1] if lang == "TH" else f[0]), T(f[3] if lang == "TH" else f[2]), T(ex)])
        table(doc, [lb["column"], lb["desc"], lb["example"]], rows, r, widths=[4.0, 8.0, 4.5])
        n += 1
        # sample entries in landscape section
        ncol = len(c["fields"])
        land = ncol > 6
        if land:
            ns = doc.add_section()
            ns.orientation = WD_ORIENT.LANDSCAPE
            ns.page_width, ns.page_height = Cm(29.7), Cm(21.0)
            ns.left_margin = ns.right_margin = Cm(1.5)
        heading(doc, f"4.{n} {lb['entries']}", 2, r)
        box(doc, lb["req"], AUTO_REQ["entries"][li], r) if not land else None
        hdrs = [T(f[1] if lang == "TH" else f[0]) for f in c["fields"]]
        if mode == "cmc":
            erows = c.get("rows") or []
        else:
            ex_row = []
            for f in c["fields"]:
                ex_row.append(T(ex_val(f, li)))
            erows = [ex_row, [""] * ncol, [""] * ncol]
        avail = 26.7 if land else 16.5
        table(doc, hdrs, erows, r, widths=[avail / ncol] * ncol, size=r.tiny if ncol > 5 else r.small)
        if land:
            ns2 = doc.add_section()
            ns2.orientation = WD_ORIENT.PORTRAIT
            ns2.page_width, ns2.page_height = Cm(21.0), Cm(29.7)
            ns2.left_margin = ns2.right_margin = Cm(2.2)
        if d["type"] == "REG":
            n += 1
            heading(doc, f"4.{n} {lb['changelog']}", 2, r)
            box(doc, lb["req"], AUTO_REQ["changelog"][li], r)
            if mode == "cmc":
                cl = c.get("changelog") or [
                    ["1", "31 March 2026", "Q1 2026 review completed; entries updated and stale entries closed.", own],
                    ["2", "30 June 2026", "Q2 2026 review completed; new entries added from audit and incident follow-up.", own],
                    ["3", "30 September 2026", "Q3 2026 review completed; owners confirmed for all open entries.", own]]
            else:
                cl = [["1", "[DD Month YYYY]" if lang == "EN" else "[วว เดือน ปปปป]",
                       "[Entries added, changed, or closed; review completed]" if lang == "EN" else "[รายการที่เพิ่ม แก้ไข หรือปิด หรือผลการทบทวน]",
                       "[Name, Role]" if lang == "EN" else "[ชื่อ ตำแหน่ง]"]]
            table(doc, ["No." if lang == "EN" else "ลำดับ", lb["date"], lb["change"], lb["author"]], cl, r, widths=[1.2, 3.0, 8.8, 3.5])

    # auto compliance / records
    if d["type"] in ("POL", "STD"):
        sect(lb["compliance"], AUTO_REQ["compliance"][li], T(compliance_text(d, lang)))
    if d["type"] in ("PRC", "PLN", "PLB"):
        sect(lb["records"], AUTO_REQ["records"][li], T(records_text(d, lang)))

    # PART B
    heading(doc, "5. " + lb["partB"], 1, r)
    for b in c.get("bp", []):
        bullet(doc, T(b[li] if mode == "tpl" else b[0]), r)
    if c.get("bp_cmc") and mode == "cmc":
        para(doc, "How CMC applies this:", r, bold=True, space_after=2)
        for b in c["bp_cmc"]:
            bullet(doc, b, r)
    if c.get("mistakes"):
        heading(doc, lb["mistakes"], 2, r)
        for b in c["mistakes"]:
            bullet(doc, T(b[li] if mode == "tpl" else b[0]), r)

    # ANNEX
    heading(doc, "6. " + lb["annex"], 1, r)
    tr = [(lb["csf"], "\n".join(f"{x}: {SUBCATS_TH[x] if lang == 'TH' else SUBCATS[x]}" for x in d["csf"])),
          (lb["iso"], ", ".join((("ข้อ " if lang == "TH" else "Clause ") + x) if not x.startswith("A.") else ("Annex " + x) for x in d["iso"])),
          (lb["nist"], ", ".join(d["nist"])),
          (lb["refs"], th_basis(d["basis"]) if lang == "TH" else d["basis"]),
          (lb["parents"], "\n".join(f"{p}: {doc_name(p, lang)}" for p in d["parents"]) or lb["none"]),
          (lb["children"], "\n".join(f"{x}: {doc_name(x, lang)}" for x in d["children"]) or lb["none"])]
    kv_table(doc, tr, r)

    fname = f"{d['id']}_{slug(d['en'])}.docx"
    path = os.path.join(outdir, fname)
    doc.save(path)
    return path

# Thai basis: keep law names readable for Thai readers
TH_BASIS = [
 ("CSA s.", "พ.ร.บ.การรักษาความมั่นคงปลอดภัยไซเบอร์ พ.ศ. 2562 มาตรา "),
 ("PDPA s.", "พ.ร.บ.คุ้มครองข้อมูลส่วนบุคคล พ.ศ. 2562 มาตรา "),
 ("Computer Crime Act B.E. 2550 s.", "พ.ร.บ.ว่าด้วยการกระทำความผิดเกี่ยวกับคอมพิวเตอร์ พ.ศ. 2550 มาตรา "),
 ("PDPC Notification on Security Measures B.E. 2565", "ประกาศคณะกรรมการคุ้มครองข้อมูลส่วนบุคคล เรื่อง มาตรการรักษาความมั่นคงปลอดภัยของผู้ควบคุมข้อมูลส่วนบุคคล พ.ศ. 2565"),
 ("PDPC Notification on Breach Notification Criteria B.E. 2565", "ประกาศคณะกรรมการคุ้มครองข้อมูลส่วนบุคคล เรื่อง หลักเกณฑ์และวิธีการในการแจ้งเหตุการละเมิดข้อมูลส่วนบุคคล พ.ศ. 2565"),
 ("CRC Notification on CII duties 2024", "ประกาศคณะกรรมการกำกับดูแลด้านความมั่นคงปลอดภัยไซเบอร์ว่าด้วยหน้าที่ของหน่วยงานโครงสร้างพื้นฐานสำคัญทางสารสนเทศ (2567)"),
 ("NCSC Notification on Code of Practice and Standard Framework B.E. 2564", "ประกาศคณะกรรมการการรักษาความมั่นคงปลอดภัยไซเบอร์แห่งชาติ เรื่อง ประมวลแนวทางปฏิบัติและกรอบมาตรฐานด้านการรักษาความมั่นคงปลอดภัยไซเบอร์ พ.ศ. 2564"),
 ("CRC Notification on cloud security for CII", "ประกาศคณะกรรมการกำกับดูแลด้านความมั่นคงปลอดภัยไซเบอร์ เรื่อง มาตรฐานขั้นต่ำด้านความมั่นคงปลอดภัยของบริการคลาวด์สำหรับ CII"),
 ("Not explicitly mandated by the PDPA; supports the risk-based measures in PDPA s.37(1) [VERIFY current PDPC guidance]",
  "PDPA ไม่ได้บังคับโดยตรง แต่สนับสนุนมาตรการตามความเสี่ยงตาม พ.ร.บ.คุ้มครองข้อมูลส่วนบุคคล พ.ศ. 2562 มาตรา 37(1) [VERIFY แนวทางของ สคส. ฉบับปัจจุบัน]"),
 ("Supports PDPA s.37(1)", "สนับสนุน พ.ร.บ.คุ้มครองข้อมูลส่วนบุคคล พ.ศ. 2562 มาตรา 37(1)"),
 ("Common practice", "แนวปฏิบัติทั่วไป"),
 ("(not mandated as documented information)", "(ไม่ได้บังคับให้จัดทำเป็นเอกสาร)"),
 ("(documented information)", "(ต้องจัดทำเป็นเอกสาร)"),
 (" and ", " และ "), (" to ", " ถึง "),
]
from basis_th import BASIS_TH
from subcats_th import SUBCATS_TH
_BT = {DOC[k]["basis"]: v for k, v in BASIS_TH.items()}

def th_basis(s):
    if s in _BT:
        return _BT[s]
    for a, b in TH_BASIS:
        s = s.replace(a, b)
    s = re.sub(r"\bs\.(\d)", r"มาตรา \1", s)
    return re.sub(r"\bClause\b", "ข้อ", s)

from register import REG
ESSENTIAL = {k for k, v in REG["documents"].items() if v["package"] == "Essential"}

def theme_text(d, lang):
    from enrich import ENRICH, THEMES, TAGS
    e = ENRICH[d["id"]]; th = THEMES[e["theme"]]
    if lang == "TH":
        return f"{th[1]} ({e['theme']}, {th[2]}); Tag: " + ", ".join(TAGS[t][1] for t in e["tags"])
    return f"{e['theme']} ({th[2]}); tags: " + ", ".join(TAGS[t][0] for t in e["tags"])

def pkg_of(d):
    if d["id"] in ESSENTIAL:
        return "Essential"
    return d["pkg"]

def transform(s, lang, mode):
    if mode == "cmc":
        return role_text(cmc_text(s), "EN", "cmc")
    return role_text(s, lang, "tpl")


TH_EX = {"01/07/2026 to 31/12/2026": "01/07/2026 ถึง 31/12/2026", "3 x 3 = 9 (Medium)": "3 x 3 = 9 (ปานกลาง)", "3 x 4 = 12 High": "3 x 4 = 12 สูง",
 "CHG-2026-0001, Normal": "CHG-2026-0001, ปกติ", "CSA s.54": "พ.ร.บ.ไซเบอร์ฯ มาตรา 54", "Compliant": "ปฏิบัติตามครบ", "External": "ภายนอก",
 "Implemented": "ดำเนินการแล้ว", "In treatment": "อยู่ระหว่างจัดการ", "Medium": "ปานกลาง", "Mitigate": "ลดความเสี่ยง", "Open": "เปิด",
 "PDPA s.37(4)": "PDPA มาตรา 37(4)", "PR-STD-007 clause 4.3": "PR-STD-007 ข้อ 4.3", "RTP-01, target Medium": "RTP-01 เป้าหมายระดับปานกลาง",
 "Yes": "ใช่", "A.5.15 Access control": "A.5.15 การควบคุมการเข้าถึง"}
def ex_val(f, li):
    ex = f[4] if len(f) > 4 else ""
    if isinstance(ex, tuple):
        return ex[li]
    if li == 1:
        return TH_EX.get(ex, ex)
    return ex

SETS = [("01_Template_TH", "TH", "tpl"), ("02_Template_EN", "EN", "tpl"), ("03_CMC_EN", "EN", "cmc")]

def folder_for(d, setname):
    fn = d["fn"]
    order = ["GV", "ID", "PR", "DE", "RS", "RC"].index(fn) + 1
    f1 = f"{order:02d}_{fn}_{FUNCTIONS[fn][0]}"
    f2 = f"{d['cat']}_{slug(CATEGORIES[d['cat']][0])}"
    p = os.path.join(OUT, setname, f1, f2)
    os.makedirs(p, exist_ok=True)
    return p

def main(only=None):
    C = load_content()
    done = 0
    for idx, d in enumerate(DOCS):
        if d["id"] not in C:
            continue
        if only and d["id"] not in only:
            continue
        for setname, lang, mode in SETS:
            render(d, C[d["id"]], lang, mode, idx, folder_for(d, setname))
        done += 1
    print("rendered", done, "docs x 3 sets; missing content:", [d["id"] for d in DOCS if d["id"] not in C][:200])

if __name__ == "__main__":
    main(sys.argv[1:] or None)
